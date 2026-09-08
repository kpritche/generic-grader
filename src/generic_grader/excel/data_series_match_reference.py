"""Test that a reference data series matches in-range or somewhere in submission."""

import unittest

from openpyxl.utils.cell import range_boundaries
from parameterized import parameterized

from generic_grader.excel._series import (
    _series_ratio,
    extract_series_from_range,
    find_best_series_window,
    find_exact_series_window,
    find_series_windows,
)
from generic_grader.excel._workbook import (
    UNCACHED_VALUES_HINT,
    formulas_lack_cached_values,
    is_formula_value,
    load_sheet,
    resolve_reference_file,
    resolve_sheet_and_range,
    resolve_submission_file,
)
from generic_grader.utils.decorators import merge_subtests, weighted
from generic_grader.utils.docs import get_wrapper
from generic_grader.utils.options import options_to_params


def doc_func(func, num, param):
    """Return parameterized docstring for data series match checks."""

    o = param.args[0]
    sheet = o.kwargs.get("sheet") or o.sheet or "<first worksheet>"
    start_cell, end_cell = o.entries
    require_formulas = o.series_require_formulas
    range_matches_reference = o.range_matches_reference
    formula_suffix = " and uses formulas" if require_formulas else ""
    if o.series_canonicalize_units:
        formula_suffix += ", comparing units by meaning rather than spelling"
    if range_matches_reference:
        # With relocation enabled a shifted-but-intact series can still earn
        # credit, so the title must not promise a same-cell requirement.
        location_clause = (
            "matches the reference values"
            if o.series_allow_relocation
            else "matches the reference values at the same cell locations"
        )
    else:
        location_clause = "matches somewhere in the submission workbook"
    return (
        f"Check that the data series in range {start_cell}:{end_cell}"
        f" on sheet `{sheet}` exactly {location_clause}"
        f"{formula_suffix}."
    )


# Relocation needs enough values that finding the series elsewhere is
# meaningful evidence; short series match too easily by coincidence.
MIN_RELOCATION_LENGTH = 3


def _iter_series_coordinates(match: dict):
    min_col, min_row, max_col, max_row = range_boundaries(
        f"{match['start_cell']}:{match['end_cell']}"
    )

    if match["orientation"] == "row":
        for col in range(min_col, max_col + 1):
            yield min_row, col
        return

    for row in range(min_row, max_row + 1):
        yield row, min_col


def build(the_options):
    """Create a class for data series matching checks."""

    the_params = options_to_params(the_options)

    class TestDataSeriesMatchReference(unittest.TestCase):
        """A class for data series matching tests."""

        wrapper = get_wrapper()

        @parameterized.expand(the_params, doc_func=doc_func)
        @merge_subtests()
        @weighted
        def test_data_series_match_reference(self, options):
            """Check that reference data series matches in the configured mode."""

            o = options
            sheet, start_cell, end_cell = resolve_sheet_and_range(o)

            sub_file = resolve_submission_file(o)
            reference_file = resolve_reference_file(o)

            ref_sheet = load_sheet(reference_file, sheet, data_only=True)
            sub_sheet = load_sheet(sub_file, sheet, data_only=True)
            sheet = ref_sheet.title

            series = extract_series_from_range(ref_sheet, start_cell, end_cell)
            search_orientation = o.kwargs.get("search_orientation", "either")
            require_formulas = o.series_require_formulas
            canonicalize_units = o.series_canonicalize_units
            range_matches_reference = o.range_matches_reference

            if require_formulas and all(value is None for value in series["values"]):
                ref_sheet_formulas = load_sheet(reference_file, sheet, data_only=False)
                series = extract_series_from_range(ref_sheet_formulas, start_cell, end_cell)

            if range_matches_reference:
                if search_orientation in {"row", "rows", "horizontal"} and series["orientation"] == "column":
                    self.fail(
                        "\n\nHint:\n"
                        + self.wrapper.fill(
                            "`kwargs['search_orientation']` is set to row-only,"
                            " but the reference `entries` define a column series."
                            + (o.hint and f"  {o.hint}" or "")
                        )
                    )
                if search_orientation in {"column", "columns", "col", "vertical"} and series["orientation"] == "row":
                    self.fail(
                        "\n\nHint:\n"
                        + self.wrapper.fill(
                            "`kwargs['search_orientation']` is set to column-only,"
                            " but the reference `entries` define a row series."
                            + (o.hint and f"  {o.hint}" or "")
                        )
                    )

                sub_series = extract_series_from_range(sub_sheet, start_cell, end_cell)
                ratio = _series_ratio(
                    sub_series["values"],
                    series["values"],
                    o.relative_tolerance,
                    o.absolute_tolerance,
                    canonicalize_units,
                )

                sub_sheet_formulas = None
                aligned_formula_gaps = []
                if require_formulas:
                    sub_sheet_formulas = load_sheet(sub_file, sheet, data_only=False)
                    aligned_formula_gaps = [
                        (row, col)
                        for row, col in _iter_series_coordinates(series)
                        if not is_formula_value(
                            sub_sheet_formulas.cell(row, col).value
                        )
                    ]

                aligned_ok = ratio >= o.ratio and not aligned_formula_gaps

                def window_is_aligned(window):
                    return (window["start_cell"], window["end_cell"]) == (
                        series["start_cell"],
                        series["end_cell"],
                    )

                def window_has_all_formulas(window):
                    return all(
                        is_formula_value(sub_sheet_formulas.cell(row, col).value)
                        for row, col in _iter_series_coordinates(window)
                    )

                # Aligned-first, relocation-fallback: when the expected window
                # fails, look for the intact series elsewhere on the sheet
                # (same orientation) and award credit if it satisfies the same
                # requirements.  Common cause: an inserted or deleted row
                # shifting an otherwise-correct block off its template cells.
                relocated_window = None
                relocation_windows = []
                allow_relocation = (
                    o.series_allow_relocation
                    and not aligned_ok
                    and len(series["values"]) >= MIN_RELOCATION_LENGTH
                )
                if allow_relocation:
                    relocation_windows = [
                        window
                        for window in find_series_windows(
                            sub_sheet,
                            series["values"],
                            series["orientation"],
                            o.relative_tolerance,
                            o.absolute_tolerance,
                            o.ratio,
                            canonicalize_units,
                        )
                        if not window_is_aligned(window)
                    ]
                    for window in relocation_windows:
                        if not require_formulas or window_has_all_formulas(window):
                            relocated_window = window
                            break

                if not aligned_ok and relocated_window is None:
                    # A workbook saved without its cached formula results
                    # reads as empty everywhere.  Say so, rather than
                    # blaming values the student cannot see are missing.
                    if ratio < o.ratio and formulas_lack_cached_values(
                        sub_file, sheet, list(_iter_series_coordinates(series))
                    ):
                        self.fail(
                            "\n\nHint:\n"
                            + self.wrapper.fill(
                                UNCACHED_VALUES_HINT
                                + (o.hint and f"  {o.hint}" or "")
                            )
                        )

                    diagnostic = ""
                    if allow_relocation and ratio < o.ratio:
                        if relocation_windows:
                            found = relocation_windows[0]
                            diagnostic = (
                                "  A matching series was found at"
                                f" `{found['start_cell']}:{found['end_cell']}`,"
                                " but every cell in it must contain a formula."
                            )
                        else:
                            best = find_best_series_window(
                                sub_sheet,
                                series["values"],
                                series["orientation"],
                                o.relative_tolerance,
                                o.absolute_tolerance,
                                canonicalize_units,
                            )
                            if (
                                best is not None
                                and best["ratio"] > ratio
                                and not window_is_aligned(best)
                            ):
                                diagnostic = (
                                    "  The closest match found was at"
                                    f" `{best['start_cell']}:{best['end_cell']}`"
                                    f" (ratio {best['ratio']:.2f})."
                                    " If your values moved, check for inserted"
                                    " or deleted rows or columns."
                                )

                    self.assertGreaterEqual(
                        ratio,
                        o.ratio,
                        msg=(
                            "\n\nHint:\n"
                            + self.wrapper.fill(
                                f"Data series in `{start_cell}:{end_cell}` on sheet `{sheet}`"
                                " did not meet the expected match ratio"
                                f" ({ratio:.2f} < {o.ratio:.2f})."
                                + diagnostic
                                + (o.hint and f"  {o.hint}" or "")
                            )
                        ),
                    )

                    for row, col in _iter_series_coordinates(series):
                        with self.subTest(row=row, column=col):
                            cell = sub_sheet_formulas.cell(row, col)
                            is_formula = is_formula_value(cell.value)
                            self.assertTrue(
                                is_formula,
                                msg=(
                                    "\n\nHint:\n"
                                    + self.wrapper.fill(
                                        f"Cell `{cell.coordinate}` on sheet `{sheet}`"
                                        " must contain a formula."
                                        + (o.hint and f"  {o.hint}" or "")
                                    )
                                ),
                            )
            else:
                exact_match = find_exact_series_window(
                    sub_sheet,
                    series["values"],
                    search_orientation,
                    o.relative_tolerance,
                    o.absolute_tolerance,
                    canonicalize_units,
                )

                if exact_match is None:
                    best_match = find_best_series_window(
                        sub_sheet,
                        series["values"],
                        search_orientation,
                        o.relative_tolerance,
                        o.absolute_tolerance,
                        canonicalize_units,
                    )

                    # Same diagnosis as the aligned branch, but the series
                    # could be anywhere, so the whole sheet is scanned.
                    # Missing an exact match is not itself a failure -- a
                    # check with a ratio below 1.0 still awards credit for
                    # a close enough window -- so this only speaks up when
                    # the check is otherwise about to fail.
                    about_to_fail = (
                        best_match is None
                        or o.ratio >= 1.0
                        or best_match["ratio"] < o.ratio
                    )
                    if about_to_fail and formulas_lack_cached_values(sub_file, sheet):
                        self.fail(
                            "\n\nHint:\n"
                            + self.wrapper.fill(
                                UNCACHED_VALUES_HINT
                                + (o.hint and f"  {o.hint}" or "")
                            )
                        )

                    if best_match is None:
                        self.fail(
                            "\n\nHint:\n"
                            + self.wrapper.fill(
                                "Could not find any candidate location for the"
                                f" series from `{series['start_cell']}:{series['end_cell']}`"
                                f" on sheet `{sheet}`."
                                + (o.hint and f"  {o.hint}" or "")
                            )
                        )

                    if o.ratio < 1.0:
                        self.assertGreaterEqual(
                            best_match["ratio"],
                            o.ratio,
                            msg=(
                                "\n\nHint:\n"
                                + self.wrapper.fill(
                                    f"Could not find data series from `{series['start_cell']}:{series['end_cell']}`"
                                    f" on sheet `{sheet}` with at least {o.ratio:.2f} match ratio."
                                    " Best match was"
                                    f" `{best_match['start_cell']}:{best_match['end_cell']}`"
                                    f" ({best_match['orientation']}) with"
                                    f" ratio {best_match['ratio']:.2f}."
                                    + (o.hint and f"  {o.hint}" or "")
                                )
                            ),
                        )
                        matched_window = best_match
                    else:
                        self.fail(
                            "\n\nHint:\n"
                            + self.wrapper.fill(
                                f"No exact data series match was found for"
                                f" `{series['start_cell']}:{series['end_cell']}` on sheet `{sheet}`."
                                " Best candidate was"
                                f" `{best_match['start_cell']}:{best_match['end_cell']}`"
                                f" ({best_match['orientation']}) with"
                                f" ratio {best_match['ratio']:.2f}."
                                + (o.hint and f"  {o.hint}" or "")
                            )
                        )

                else:
                    matched_window = exact_match

                if require_formulas:
                    sub_sheet_formulas = load_sheet(sub_file, sheet, data_only=False)
                    for row, col in _iter_series_coordinates(matched_window):
                        with self.subTest(row=row, column=col):
                            cell = sub_sheet_formulas.cell(row, col)
                            is_formula = is_formula_value(cell.value)
                            self.assertTrue(
                                is_formula,
                                msg=(
                                    "\n\nHint:\n"
                                    + self.wrapper.fill(
                                        f"Matched series cell `{cell.coordinate}` on sheet `{sheet}`"
                                        " is not a formula cell."
                                        + (o.hint and f"  {o.hint}" or "")
                                    )
                                ),
                            )

            self.set_score(self, o.weight)

    return TestDataSeriesMatchReference
